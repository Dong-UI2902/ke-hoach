from openpyxl import load_workbook

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
workbook = load_workbook(path, read_only=False, data_only=False)
print("ACTIVE", workbook.index(workbook.active), workbook.active.title.encode("unicode_escape").decode())

for sheet in workbook.worksheets:
    selections = []
    if sheet.sheet_view.selection:
        for selection in sheet.sheet_view.selection:
            selections.append((selection.pane, selection.activeCell, selection.sqref))
    print(
        "VIEW",
        sheet.title.encode("unicode_escape").decode(),
        "freeze", str(sheet.freeze_panes),
        "topLeft", sheet.sheet_view.topLeftCell,
        "selections", selections,
    )

sheet = workbook["Timeline"]
print("TIMELINE_SIZE", sheet.max_row, sheet.max_column)
print("MERGES", [str(item) for item in sheet.merged_cells.ranges])
for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 30), min_col=1, max_col=min(sheet.max_column, 15)):
    values = [cell.value for cell in row]
    if any(value not in (None, "") for value in values):
        print(row[0].row, [str(value).encode("unicode_escape").decode() if value is not None else "" for value in values])
