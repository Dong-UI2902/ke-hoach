from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
sheet = load_workbook(path, read_only=False, data_only=False)["Timeline GTM"]

for row in range(3, 37):
    segments = []
    start = None
    previous = None
    color = None
    for col in range(9, 58):
        cell = sheet.cell(row, col)
        fill = cell.fill
        rgb = fill.fgColor.rgb if fill.fill_type == "solid" and fill.fgColor.type == "rgb" else None
        if rgb and rgb not in ("00000000", "FFFFFFFF"):
            if start is None:
                start = col
                previous = col
                color = rgb
            elif col == previous + 1 and rgb == color:
                previous = col
            else:
                segments.append((get_column_letter(start), get_column_letter(previous), color))
                start = col
                previous = col
                color = rgb
        elif start is not None:
            segments.append((get_column_letter(start), get_column_letter(previous), color))
            start = previous = color = None
    if start is not None:
        segments.append((get_column_letter(start), get_column_letter(previous), color))
    if segments:
        print(row, segments)
