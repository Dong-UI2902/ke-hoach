from openpyxl import load_workbook

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
workbook = load_workbook(path, read_only=False, data_only=False)

for sheet_name in ["Giá Needly đơn đầu", "Phân tích từ khóa_Đối thủ cạnh tranh", "Đối thủ"]:
    sheet = workbook[sheet_name]
    print("\nSHEET", sheet_name.encode("unicode_escape").decode(), sheet.sheet_state, sheet.max_row, sheet.max_column)
    print("MERGES", [str(r) for r in sheet.merged_cells.ranges])
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        if any(value not in (None, "") for value in values):
            rendered = [str(value).encode("unicode_escape").decode() if value is not None else "" for value in values]
            print(row[0].row, rendered)
