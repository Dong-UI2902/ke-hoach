from openpyxl import load_workbook

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
workbook = load_workbook(path, read_only=False, data_only=False)
sheet = workbook["Timeline GTM"]

assert sheet.freeze_panes == "D3"
assert sheet["A3"].value == "W1"
assert sheet["A7"].value == "W5"
assert sheet["F3"].value.strftime("%d/%m/%Y") == "30/09/2026"
assert sheet["G7"].value.strftime("%d/%m/%Y") == "27/10/2026"
assert [sheet[f"H{row}"].value for row in range(3, 8)] == [3, 5, 5, 5, 2]
assert sheet["L2"].value.strftime("%d/%m/%Y") == "30/09/2026"
assert sheet["AQ2"].value.strftime("%d/%m/%Y") == "31/10/2026"
assert len(sheet.data_validations.dataValidation) == 1

formula_count = sum(
    1 for ws in workbook.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f"
)
error_count = sum(
    1
    for ws in workbook.worksheets
    for row in ws.iter_rows()
    for cell in row
    if isinstance(cell.value, str) and cell.value.startswith("#")
)
print("FREEZE", sheet.freeze_panes)
print("TIMELINE_RANGE", sheet.calculate_dimension())
print("VALIDATIONS", len(sheet.data_validations.dataValidation))
print("FORMULAS", formula_count)
print("ERROR_STRINGS", error_count)
print("CHECK_OK")
