from openpyxl import load_workbook

paths = [
    r"C:\Users\Admin\Desktop\InviBrowser_GTM_Execution_v1.9_Clean.xlsx",
    r"C:\Users\Admin\Desktop\InviBrowser_Sales_Marketing_Dashboard_v2.0_Clean.xlsx",
]

for path in paths:
    workbook = load_workbook(path, read_only=False, data_only=False)
    print("\nWORKBOOK", path.encode("unicode_escape").decode())
    for sheet in workbook.worksheets:
        nonempty = []
        formulas = 0
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    nonempty.append(cell)
                if cell.data_type == "f":
                    formulas += 1
        if nonempty:
            bounds = (
                min(c.row for c in nonempty),
                min(c.column for c in nonempty),
                max(c.row for c in nonempty),
                max(c.column for c in nonempty),
            )
        else:
            bounds = None
        print(
            "SHEET",
            sheet.title.encode("unicode_escape").decode(),
            "STATE", sheet.sheet_state,
            "BOUNDS", bounds,
            "NONEMPTY", len(nonempty),
            "FORMULAS", formulas,
        )
