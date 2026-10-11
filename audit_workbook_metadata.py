from collections import Counter
from openpyxl import load_workbook

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
workbook = load_workbook(path, read_only=False, data_only=False)

print("WORKBOOK_ACTIVE", workbook.active.title.encode("unicode_escape").decode())
print("SHEET_COUNT", len(workbook.worksheets))

needly_hits = []
error_hits = []
formula_hits = []
duplicate_text = Counter()

for sheet in workbook.worksheets:
    nonempty = []
    for row in sheet.iter_rows():
        for cell in row:
            value = cell.value
            if value not in (None, ""):
                nonempty.append(cell)
                if isinstance(value, str):
                    duplicate_text[value.strip()] += 1
                    if "needly" in value.lower():
                        needly_hits.append((sheet.title, cell.coordinate, value))
                    if value.startswith("#"):
                        error_hits.append((sheet.title, cell.coordinate, value))
                if cell.data_type == "f":
                    formula_hits.append((sheet.title, cell.coordinate, value))
    if nonempty:
        min_row = min(cell.row for cell in nonempty)
        max_row = max(cell.row for cell in nonempty)
        min_col = min(cell.column for cell in nonempty)
        max_col = max(cell.column for cell in nonempty)
        bounds = (min_row, min_col, max_row, max_col)
    else:
        bounds = None
    hidden_rows = [index for index, dimension in sheet.row_dimensions.items() if dimension.hidden]
    hidden_cols = [key for key, dimension in sheet.column_dimensions.items() if dimension.hidden]
    validations = sum(len(validation.sqref.ranges) for validation in sheet.data_validations.dataValidation)
    print("\nSHEET", sheet.title.encode("unicode_escape").decode())
    print("STATE", sheet.sheet_state, "NONEMPTY", len(nonempty), "BOUNDS", bounds, "MAX", sheet.max_row, sheet.max_column)
    print("FREEZE", sheet.freeze_panes, "TOPLEFT", sheet.sheet_view.topLeftCell)
    print("MERGES", len(sheet.merged_cells.ranges), "VALIDATION_RANGES", validations)
    print("HIDDEN_ROWS", hidden_rows[:30], "COUNT", len(hidden_rows))
    print("HIDDEN_COLS", hidden_cols, "COUNT", len(hidden_cols))
    print("AUTO_FILTER", sheet.auto_filter.ref, "PRINT_AREA", sheet.print_area)

print("\nFORMULA_COUNT", len(formula_hits))
print("ERROR_COUNT", len(error_hits))
print("NEEDLY_COUNT", len(needly_hits))
for item in needly_hits:
    print("NEEDLY", item[0].encode("unicode_escape").decode(), item[1], str(item[2]).encode("unicode_escape").decode())
