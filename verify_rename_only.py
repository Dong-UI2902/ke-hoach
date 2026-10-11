from openpyxl import load_workbook

before_path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.pre_sheet_rename_backup.xlsx"
after_path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"

before = load_workbook(before_path, read_only=False, data_only=False)
after = load_workbook(after_path, read_only=False, data_only=False)

mapping = {
    "Kế hoạch kinh doanh": "Kế hoạch kinh doanh",
    "Tài chính InviBrowser": "Kế hoạch tài chính Needly",
    "Khung tài chính mẫu": "Kế hoạch tài chính Needly（Copy）",
    "Giá InviBrowser": "Giá InviBrowser",
    "Nhân sự dự án": "Nhân sự dự án",
    "Đối thủ & nội dung": "Đối thủ",
    "Timeline GTM": "Timeline",
}

problems = []
for new_name, old_name in mapping.items():
    new_sheet = after[new_name]
    old_sheet = before[old_name]
    max_row = max(new_sheet.max_row, old_sheet.max_row)
    max_col = max(new_sheet.max_column, old_sheet.max_column)
    for row in range(1, max_row + 1):
        for col in range(1, max_col + 1):
            new_cell = new_sheet.cell(row, col)
            old_cell = old_sheet.cell(row, col)
            if new_cell.value != old_cell.value:
                problems.append((new_name, new_cell.coordinate, "value"))
            if new_cell.style_id != old_cell.style_id:
                problems.append((new_name, new_cell.coordinate, "style"))
    if {str(r) for r in new_sheet.merged_cells.ranges} != {str(r) for r in old_sheet.merged_cells.ranges}:
        problems.append((new_name, "-", "merges"))
    if str(new_sheet.freeze_panes) != str(old_sheet.freeze_panes):
        problems.append((new_name, "-", "freeze"))
    if len(new_sheet.data_validations.dataValidation) != len(old_sheet.data_validations.dataValidation):
        problems.append((new_name, "-", "data_validations"))
    if len(new_sheet.conditional_formatting) != len(old_sheet.conditional_formatting):
        problems.append((new_name, "-", "conditional_formatting"))

print("DIFF_COUNT", len(problems))
for problem in problems[:50]:
    print(problem)
