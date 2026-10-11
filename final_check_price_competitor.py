from openpyxl import load_workbook

path = r"D:\kehoach\InviBrowser_Business_Plan_Giu_Nguyen_Format_v2.xlsx"
workbook = load_workbook(path, read_only=False, data_only=False)
names = [sheet.title for sheet in workbook.worksheets]

assert "Giá InviBrowser" in names
assert "Giá Needly đơn đầu" not in names
assert "Phân tích từ khóa_Đối thủ cạnh tranh" not in names
assert "KH doanh thu và P&L" not in names
assert workbook["Giá InviBrowser"]["C3"].value == "Solo Pro"
assert workbook["Giá InviBrowser"]["F4"].value == "520.000đ/tháng"
assert workbook["Đối thủ"]["A34"].value == "4. CHỦ ĐỀ TÌM KIẾM & NỘI DUNG CẠNH TRANH"

formula_count = sum(
    1
    for sheet in workbook.worksheets
    for row in sheet.iter_rows()
    for cell in row
    if cell.data_type == "f"
)
error_strings = sum(
    1
    for sheet in workbook.worksheets
    for row in sheet.iter_rows()
    for cell in row
    if isinstance(cell.value, str) and cell.value.startswith("#")
)
reference = workbook["Kế hoạch tài chính Needly（Copy）"]
reference_month_values = sum(
    1
    for row in reference.iter_rows(min_col=5, max_col=23)
    for cell in row
    if cell.value not in (None, "")
)

print("SHEETS", [name.encode("unicode_escape").decode() for name in names])
print("FORMULAS", formula_count)
print("ERROR_STRINGS", error_strings)
print("REFERENCE_E_W_NONEMPTY", reference_month_values)
print("CHECK_OK")
